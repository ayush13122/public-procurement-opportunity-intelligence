"""Prepare a structured Excel workbook for reviewing procurement niche mappings.

Run from the project root:
    .venv\\Scripts\\python.exe src\\prepare_taxonomy_signoff.py

This script reads existing reports and data only. It does not modify raw or cleaned data.
"""
from __future__ import annotations

from pathlib import Path
import sys
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORTS_DIR = PROJECT_ROOT / "reports"
TAXONOMY_DIR = REPORTS_DIR / "taxonomy_review"
CLEANED_DATA = PROJECT_ROOT / "data" / "processed" / "assam_procurement_cleaned.csv"
OUTPUT_DIR = TAXONOMY_DIR
OUTPUT_FILE = OUTPUT_DIR / "Taxonomy_Review_Workbook.xlsx"

REQUIRED = {
    "classification_review": REPORTS_DIR / "assam_niche_mapping_review.csv",
    "assignments": REPORTS_DIR / "assam_record_niche_assignments.csv",
    "review_queue": TAXONOMY_DIR / "00_classification_review_queue.csv",
    "flagged_records": TAXONOMY_DIR / "01_records_requiring_review.csv",
    "cleaned_data": CLEANED_DATA,
}

DECISION_OPTIONS = [
    "Approve",
    "Change niche",
    "Keep Other / Needs Review",
    "Needs tender-document review",
]


def load_csv(path: Path, label: str) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Missing {label}: {path}")
    return pd.read_csv(path, low_memory=False)


def first_titles_by_classification(cleaned: pd.DataFrame, title_limit: int = 3) -> pd.DataFrame:
    required = {"tenderclassification/description", "tender/title"}
    if not required.issubset(cleaned.columns):
        return pd.DataFrame(columns=["original_classification", "sample_tender_titles"])

    rows = []
    for label, group in cleaned.groupby("tenderclassification/description", dropna=False, sort=False):
        titles = group["tender/title"].dropna().astype(str).drop_duplicates().head(title_limit).tolist()
        rows.append({
            "original_classification": "(MISSING)" if pd.isna(label) else str(label),
            "sample_tender_titles": "\n---\n".join(titles),
        })
    return pd.DataFrame(rows)


def ensure_review_columns(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for column in ("reviewer_decision", "approved_analytical_niche", "reviewer_notes"):
        if column not in result.columns:
            result[column] = ""
    return result


def create_samples(flagged: pd.DataFrame, per_niche: int = 8, per_class_unresolved: int = 5) -> tuple[pd.DataFrame, pd.DataFrame]:
    candidates = flagged[flagged.get("mapping_method", pd.Series(index=flagged.index, dtype=object)).eq("TITLE_KEYWORD_CANDIDATE")].copy()
    candidate_parts = []
    if not candidates.empty and "suggested_analytical_niche" in candidates.columns:
        for _, group in candidates.groupby("suggested_analytical_niche", dropna=False, sort=True):
            candidate_parts.append(group.sample(n=min(per_niche, len(group)), random_state=42))
    candidate_sample = pd.concat(candidate_parts, ignore_index=True) if candidate_parts else candidates.head(0).copy()

    unresolved = flagged[flagged.get("mapping_method", pd.Series(index=flagged.index, dtype=object)).eq("UNRESOLVED")].copy()
    unresolved_parts = []
    class_col = "tenderclassification/description"
    if not unresolved.empty and class_col in unresolved.columns:
        for _, group in unresolved.groupby(class_col, dropna=False, sort=True):
            unresolved_parts.append(group.sample(n=min(per_class_unresolved, len(group)), random_state=43))
    elif not unresolved.empty:
        unresolved_parts.append(unresolved.sample(n=min(30, len(unresolved)), random_state=43))
    unresolved_sample = pd.concat(unresolved_parts, ignore_index=True) if unresolved_parts else unresolved.head(0).copy()

    for df in (candidate_sample, unresolved_sample):
        for column in ("reviewer_decision", "approved_analytical_niche", "reviewer_notes"):
            if column not in df.columns:
                df[column] = ""
    return candidate_sample, unresolved_sample


def make_niche_counts(assignments: pd.DataFrame) -> pd.DataFrame:
    niche_col = "suggested_analytical_niche"
    method_col = "mapping_method"
    if niche_col not in assignments.columns:
        return pd.DataFrame()
    counts = assignments.groupby(niche_col, dropna=False).agg(
        records=(niche_col, "size"),
        unique_buyers=("buyer/name", "nunique") if "buyer/name" in assignments.columns else (niche_col, "size"),
        source_classifications=("tenderclassification/description", "nunique") if "tenderclassification/description" in assignments.columns else (niche_col, "size"),
    ).reset_index().rename(columns={niche_col: "suggested_analytical_niche"})
    counts["record_pct"] = (counts["records"] / max(len(assignments), 1) * 100).round(2)
    if method_col in assignments.columns:
        methods = assignments.pivot_table(
            index=niche_col, columns=method_col, values="ocid" if "ocid" in assignments.columns else niche_col,
            aggfunc="count", fill_value=0, dropna=False
        ).reset_index().rename(columns={niche_col: "suggested_analytical_niche"})
        counts = counts.merge(methods, on="suggested_analytical_niche", how="left")
    return counts.sort_values("records", ascending=False)


def style_workbook(writer: pd.ExcelWriter) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 and ws.max_column > 0 else "A1"
        ws.row_dimensions[1].height = 30
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in ws.columns:
            col_idx = column_cells[0].column
            max_len = 10
            for cell in column_cells[:150]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value.splitlines()[0]) if value else 0, 55))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 55)
        # Longer text columns should be readable but not take over the sheet.
        for idx, cell in enumerate(ws[1], start=1):
            label = str(cell.value or "").lower()
            if any(term in label for term in ("title", "rationale", "notes", "reason", "description", "summary")):
                ws.column_dimensions[get_column_letter(idx)].width = 48

        header_to_col = {str(c.value): c.column for c in ws[1]}
        for field in ("reviewer_decision", "Decision"): 
            if field in header_to_col and ws.max_row >= 2:
                col_letter = get_column_letter(header_to_col[field])
                dv = DataValidation(type="list", formula1='"' + ",".join(DECISION_OPTIONS) + '"', allow_blank=True)
                dv.error = "Please select a listed review decision."
                dv.errorTitle = "Invalid decision"
                dv.prompt = "Choose a decision after reviewing the title and classification."
                dv.promptTitle = "Taxonomy review"
                ws.add_data_validation(dv)
                dv.add(f"{col_letter}2:{col_letter}{ws.max_row}")


def main() -> int:
    print("Preparing taxonomy sign-off workbook...")
    for label, path in REQUIRED.items():
        if not path.exists():
            print(f"ERROR: Required input missing ({label}): {path}", file=sys.stderr)
            print("Run category profiling, cleaning, EDA, and taxonomy review scripts first.", file=sys.stderr)
            return 2

    classification_review = load_csv(REQUIRED["classification_review"], "classification review")
    assignments = load_csv(REQUIRED["assignments"], "record-level niche assignments")
    queue = load_csv(REQUIRED["review_queue"], "classification review queue")
    flagged = load_csv(REQUIRED["flagged_records"], "flagged records")
    cleaned = load_csv(REQUIRED["cleaned_data"], "cleaned dataset")

    title_samples = first_titles_by_classification(cleaned)
    queue = queue.merge(title_samples, on="original_classification", how="left")
    queue = ensure_review_columns(queue)

    candidate_sample, unresolved_sample = create_samples(flagged)
    flagged_with_review = ensure_review_columns(flagged)
    niche_counts = make_niche_counts(assignments)

    method_summary = assignments.groupby("mapping_method", dropna=False).agg(
        records=("mapping_method", "size"),
        unique_buyers=("buyer/name", "nunique") if "buyer/name" in assignments.columns else ("mapping_method", "size"),
        flagged_records=("manual_review_required", lambda s: (s.astype(str).str.upper().eq("YES")).sum()) if "manual_review_required" in assignments.columns else ("mapping_method", "size"),
    ).reset_index()
    method_summary["record_pct"] = (method_summary["records"] / max(len(assignments), 1) * 100).round(2)

    readme = pd.DataFrame([
        {"Step": "1. Review source labels", "Instructions": "Use Source_Label_Review. Check original classification, record count, proposed niche, rationale and sample titles. Fill reviewer_decision, approved_analytical_niche and reviewer_notes."},
        {"Step": "2. Review keyword matches", "Instructions": "Use Keyword_Candidate_Sample to inspect sample titles by proposed niche. Candidate mapping is not ground truth."},
        {"Step": "3. Review unresolved examples", "Instructions": "Use Unresolved_Sample to see examples that did not get a confident source/title mapping."},
        {"Step": "4. Interpret counts", "Instructions": "Niche_Counts describes current provisional assignments, not validated accuracy."},
        {"Step": "5. Preserve source data", "Instructions": "Do not edit the raw ZIP or cleaned CSV in Excel. This workbook is a review artifact only."},
        {"Step": "Important limitation", "Instructions": "The workbook samples records to reduce manual effort. Review a representative sample and high-volume/ambiguous labels before approving taxonomy rules."},
    ])

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="START_HERE", index=False)
        queue.to_excel(writer, sheet_name="Source_Label_Review", index=False)
        candidate_sample.to_excel(writer, sheet_name="Keyword_Candidate_Sample", index=False)
        unresolved_sample.to_excel(writer, sheet_name="Unresolved_Sample", index=False)
        flagged_with_review.to_excel(writer, sheet_name="All_Flagged_Records", index=False)
        niche_counts.to_excel(writer, sheet_name="Niche_Counts", index=False)
        method_summary.to_excel(writer, sheet_name="Mapping_Methods", index=False)
        style_workbook(writer)

    summary_path = OUTPUT_DIR / "TAXONOMY_SIGNOFF_WORKBOOK_SUMMARY.md"
    candidate_count = int(flagged.get("mapping_method", pd.Series(dtype=object)).eq("TITLE_KEYWORD_CANDIDATE").sum())
    unresolved_count = int(flagged.get("mapping_method", pd.Series(dtype=object)).eq("UNRESOLVED").sum())
    summary_path.write_text(
        "# Taxonomy Sign-off Workbook Summary\n\n"
        f"- Cleaned data records: {len(cleaned):,}\n"
        f"- Record-level assignments: {len(assignments):,}\n"
        f"- Source classification labels in queue: {len(queue):,}\n"
        f"- Flagged records: {len(flagged):,}\n"
        f"- Title-keyword candidates: {candidate_count:,}\n"
        f"- Unresolved records: {unresolved_count:,}\n"
        f"- Candidate review sample rows: {len(candidate_sample):,}\n"
        f"- Unresolved review sample rows: {len(unresolved_sample):,}\n\n"
        "The workbook contains provisional mappings. It does not claim classification accuracy. Fill reviewer decisions before generating final approved assignments.\n",
        encoding="utf-8",
    )

    print("Taxonomy sign-off workbook generated.")
    print(f"Cleaned rows: {len(cleaned):,}")
    print(f"Assignment rows: {len(assignments):,}")
    print(f"Source labels: {len(queue):,}")
    print(f"Flagged records: {len(flagged):,} (keyword candidates: {candidate_count:,}; unresolved: {unresolved_count:,})")
    print(f"Candidate sample rows: {len(candidate_sample):,}; unresolved sample rows: {len(unresolved_sample):,}")
    print(f"Workbook: {OUTPUT_FILE}")
    print(f"Summary: {summary_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
